from rest_framework import viewsets, permissions, status, filters, serializers
from rest_framework.decorators import action
from rest_framework.response import Response
from django.db.models import Q, Count
from django_filters.rest_framework import DjangoFilterBackend
from django.http import HttpResponse
import csv
import io
import openpyxl

from inventory.models import (
    ItemCategory, Unit, UnitConversion, Brand, InventoryItem,
    CustomFieldDefinition, InventoryLocationType, InventoryLocation,
    StockLedger, StockSummary, InventoryTransfer,
)
from inventory.serializers import (
    ItemCategorySerializer, ItemCategoryListSerializer, CategoryTreeSerializer,
    UnitSerializer, UnitListSerializer, UnitConversionSerializer,
    BrandSerializer, BrandListSerializer,
    InventoryItemSerializer, InventoryItemListSerializer,
    CustomFieldDefinitionSerializer, BulkImportSerializer,
    LocationTypeSerializer, LocationTypeListSerializer,
    LocationSerializer, LocationListSerializer, LocationTreeSerializer,
    StockLedgerSerializer, CreateLedgerEntrySerializer,
    StockAvailabilitySerializer, LowStockSerializer,
    TransferListSerializer, TransferDetailSerializer,
    CreateTransferSerializer, ReceiveTransferSerializer,
)
from inventory.permissions import (
    CanViewItems, CanCreateItems, CanEditItems,
    CanDeleteItems, CanImportItems, CanExportItems,
    CanViewCategories, CanCreateCategories, CanEditCategories, CanDeleteCategories,
    CanViewUnits, CanCreateUnits, CanEditUnits, CanDeleteUnits,
    CanViewBrands, CanCreateBrands, CanEditBrands, CanDeleteBrands,
    CanViewLocationTypes, CanCreateLocationTypes, CanEditLocationTypes, CanDeleteLocationTypes,
    CanViewLocations, CanCreateLocations, CanEditLocations, CanDeleteLocations,
    CanViewStock, CanExportStock, CanViewStockSnapshot, CanViewStockValuation,
    CanViewTransfers, CanCreateTransfers, CanEditTransfers, CanDeleteTransfers,
    CanSubmitTransfer, CanApproveTransfer, CanReceiveTransfer, CanExportTransfers,
)
from datetime import datetime
from django.utils import timezone
from inventory.services.stock_engine import (
    get_all_availability, get_item_availability, get_location_availability,
    get_low_stock_items, get_out_of_stock_items,
    get_valuation, get_snapshot, create_ledger_entry,
)
from inventory.services.transfer_service import (
    create_transfer, submit_transfer, approve_transfer,
    reject_transfer, dispatch_transfer, receive_transfer,
    cancel_transfer,
)


def _get_tenant_id(request):
    """Helper to get tenant_id from request user."""
    return request.user.organization_id


def _get_permissions_for_action(action_map, default):
    """Helper to resolve permission classes based on action."""
    def get_permissions(view):
        for action, perm_class in action_map.items():
            if view.action == action:
                return [perm_class()]
        return [default()]
    return get_permissions


# ============================================================================
# CATEGORY VIEWSET
# ============================================================================

class ItemCategoryViewSet(viewsets.ModelViewSet):
    queryset = ItemCategory.objects.all()
    filter_backends = [DjangoFilterBackend, filters.SearchFilter, filters.OrderingFilter]
    filterset_fields = {
        'status': ['exact'],
        'parent': ['exact', 'isnull'],
    }
    search_fields = ['category_code', 'category_name', 'description']
    ordering_fields = ['category_code', 'category_name', 'created_at', 'status']
    ordering = ['category_name']

    def get_serializer_class(self):
        if self.action == 'list':
            return ItemCategorySerializer
        elif self.action == 'tree':
            return CategoryTreeSerializer
        return ItemCategorySerializer

    def get_queryset(self):
        tenant_id = _get_tenant_id(self.request)
        qs = ItemCategory.objects.filter(tenant_id=tenant_id)

        # Filter archived unless explicitly requested
        show_archived = self.request.query_params.get('show_archived', 'false') == 'true'
        if not show_archived and self.action != 'tree':
            qs = qs.exclude(status='ARCHIVED')

        return qs

    def get_permissions(self):
        perms_map = {
            'create': [CanCreateCategories],
            'update': [CanEditCategories],
            'partial_update': [CanEditCategories],
            'destroy': [CanDeleteCategories],
        }
        for action, perms in perms_map.items():
            if self.action == action:
                return [p() for p in perms]
        return [CanViewCategories()]

    def perform_create(self, serializer):
        user = self.request.user
        serializer.save(
            tenant_id=_get_tenant_id(self.request),
            created_by=user,
            updated_by=user,
        )

    def perform_update(self, serializer):
        serializer.save(updated_by=self.request.user)

    def perform_destroy(self, instance):
        instance.status = 'ARCHIVED'
        instance.save()

    @action(detail=True, methods=['post'])
    def archive(self, request, pk=None):
        """Archive a category."""
        category = self.get_object()
        category.status = 'ARCHIVED'
        category.save()
        return Response({
            'success': True,
            'message': f"Category '{category.category_name}' archived."
        })

    @action(detail=True, methods=['post'])
    def restore(self, request, pk=None):
        """Restore an archived category."""
        category = self.get_object()
        category.status = 'ACTIVE'
        category.save()
        return Response({
            'success': True,
            'message': f"Category '{category.category_name}' restored."
        })

    @action(detail=False, methods=['get'])
    def tree(self, request):
        """Return hierarchical category tree with item counts."""
        tenant_id = _get_tenant_id(request)
        # Only top-level categories (no parent)
        categories = ItemCategory.objects.filter(
            tenant_id=tenant_id,
            parent__isnull=True,
        )
        show_archived = request.query_params.get('show_archived', 'false') == 'true'
        if not show_archived:
            categories = categories.exclude(status='ARCHIVED')

        serializer = CategoryTreeSerializer(
            categories, many=True,
            context={'request': request, 'filters': {'status': None}}
        )
        return Response(serializer.data)

    @action(detail=False, methods=['get'], url_path='export')
    def export(self, request):
        """Export categories to CSV."""
        tenant_id = _get_tenant_id(request)
        qs = ItemCategory.objects.filter(tenant_id=tenant_id)

        status_filter = request.query_params.get('status')
        if status_filter:
            qs = qs.filter(status=status_filter)

        format = request.query_params.get('format', 'csv').lower()
        if format == 'csv':
            return self._export_csv(qs)
        else:
            return self._export_excel(qs)

    def _export_csv(self, qs):
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(['Category Code', 'Category Name', 'Parent', 'Description', 'Status'])
        for cat in qs:
            writer.writerow([
                cat.category_code,
                cat.category_name,
                cat.parent.category_code if cat.parent else '',
                cat.description,
                cat.status,
            ])
        output.seek(0)
        response = HttpResponse(output.getvalue(), content_type='text/csv')
        response['Content-Disposition'] = 'attachment; filename="categories.csv"'
        return response

    def _export_excel(self, qs):
        """Export categories as Excel (.xlsx)."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Categories'
        ws.append(['Category Code', 'Category Name', 'Parent', 'Description', 'Status'])
        for cat in qs:
            ws.append([
                cat.category_code,
                cat.category_name,
                cat.parent.category_code if cat.parent else '',
                cat.description,
                cat.status,
            ])
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="categories.xlsx"'
        return response

    @action(detail=False, methods=['post'], url_path='import')
    def import_bulk(self, request):
        """Bulk import categories from CSV/Excel or JSON."""
        file = request.FILES.get('file')
        if file:
            return self._import_from_file(file, request)

        serializer = BulkImportSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        return self._bulk_create(serializer.validated_data['items'], request)

    def _import_from_file(self, file, request):
        tenant_id = _get_tenant_id(request)
        created = []
        errors = []
        try:
            if file.name.endswith('.csv'):
                reader = csv.DictReader(io.StringIO(file.read().decode('utf-8-sig')))
                rows = list(reader)
            elif file.name.endswith(('.xlsx', '.xls')):
                wb = openpyxl.load_workbook(file, read_only=True)
                ws = wb.active
                headers = [cell.value for cell in next(ws.iter_rows())]
                rows = []
                for row in ws.iter_rows(min_row=2, values_only=True):
                    rows.append(dict(zip(headers, row)))
            else:
                return Response({'error': 'Unsupported file format.'}, status=status.HTTP_400_BAD_REQUEST)

            for idx, row in enumerate(rows):
                code = str(row.get('Category Code', row.get('category_code', ''))).strip()
                name = str(row.get('Category Name', row.get('category_name', ''))).strip()
                if not code or not name:
                    errors.append({'index': idx, 'error': 'Category Code and Name are required.'})
                    continue

                if ItemCategory.objects.filter(tenant_id=tenant_id, category_code=code).exists():
                    errors.append({'index': idx, 'error': f"Code '{code}' already exists."})
                    continue

                parent_code = str(row.get('Parent', row.get('parent', ''))).strip()
                parent = None
                if parent_code:
                    parent = ItemCategory.objects.filter(
                        tenant_id=tenant_id, category_code=parent_code
                    ).first()

                cat = ItemCategory(
                    tenant_id=tenant_id,
                    category_code=code,
                    category_name=name,
                    parent=parent,
                    description=str(row.get('Description', row.get('description', ''))).strip(),
                    status='ACTIVE',
                )
                cat.save()
                created.append(ItemCategoryListSerializer(cat).data)

        except Exception as e:
            return Response({'error': f'Failed to parse file: {str(e)}'}, status=status.HTTP_400_BAD_REQUEST)

        return Response({
            'created_count': len(created),
            'error_count': len(errors),
            'created': created,
            'errors': errors,
        }, status=status.HTTP_201_CREATED)

    def _bulk_create(self, items_data, request):
        tenant_id = _get_tenant_id(request)
        created = []
        errors = []

        for idx, item in enumerate(items_data):
            code = item.get('category_code', '').strip()
            name = item.get('category_name', '').strip()
            if not code or not name:
                errors.append({'index': idx, 'error': 'Category Code and Name are required.'})
                continue

            if ItemCategory.objects.filter(tenant_id=tenant_id, category_code=code).exists():
                errors.append({'index': idx, 'error': f"Code '{code}' already exists."})
                continue

            cat = ItemCategory(
                tenant_id=tenant_id,
                category_code=code,
                category_name=name,
                description=item.get('description', ''),
                status='ACTIVE',
            )
            cat.save()
            created.append(ItemCategoryListSerializer(cat).data)

        return Response({
            'created_count': len(created),
            'error_count': len(errors),
            'created': created,
            'errors': errors,
        }, status=status.HTTP_201_CREATED)


# ============================================================================
# UNIT VIEWSET
# ============================================================================

class UnitViewSet(viewsets.ModelViewSet):
    filter_backends = [DjangoFilterBackend, filters.SearchFilter, filters.OrderingFilter]
    filterset_fields = {
        'status': ['exact'],
    }
    search_fields = ['unit_code', 'unit_name', 'symbol']
    ordering_fields = ['unit_code', 'unit_name', 'created_at', 'status']
    ordering = ['unit_name']

    def get_serializer_class(self):
        if self.action == 'list':
            return UnitSerializer
        return UnitSerializer

    def get_queryset(self):
        tenant_id = _get_tenant_id(self.request)
        qs = Unit.objects.filter(tenant_id=tenant_id)

        show_archived = self.request.query_params.get('show_archived', 'false') == 'true'
        if not show_archived:
            qs = qs.exclude(status='ARCHIVED')

        return qs

    def get_permissions(self):
        perms_map = {
            'create': [CanCreateUnits],
            'update': [CanEditUnits],
            'partial_update': [CanEditUnits],
            'destroy': [CanDeleteUnits],
        }
        for action, perms in perms_map.items():
            if self.action == action:
                return [p() for p in perms]
        return [CanViewUnits()]

    def perform_create(self, serializer):
        user = self.request.user
        serializer.save(
            tenant_id=_get_tenant_id(self.request),
        )

    def perform_destroy(self, instance):
        # Check if unit is used by active items
        if instance.items.filter(status='ACTIVE').exists():
            raise serializers.ValidationError(
                "Cannot archive unit used by active items."
            )
        instance.status = 'ARCHIVED'
        instance.save()

    @action(detail=True, methods=['post'])
    def archive(self, request, pk=None):
        """Archive a unit. Cannot archive if used by active items."""
        unit = self.get_object()
        if unit.items.filter(status='ACTIVE').exists():
            return Response({
                'success': False,
                'message': f"Cannot archive '{unit.unit_name}' — it is used by active items."
            }, status=status.HTTP_400_BAD_REQUEST)
        unit.status = 'ARCHIVED'
        unit.save()
        return Response({
            'success': True,
            'message': f"Unit '{unit.unit_name}' archived."
        })

    @action(detail=True, methods=['post'])
    def restore(self, request, pk=None):
        """Restore an archived unit."""
        unit = self.get_object()
        unit.status = 'ACTIVE'
        unit.save()
        return Response({
            'success': True,
            'message': f"Unit '{unit.unit_name}' restored."
        })

    @action(detail=True, methods=['get', 'post', 'delete'], url_path='conversions')
    def conversions(self, request, pk=None):
        """Manage conversions for a unit."""
        unit = self.get_object()

        if request.method == 'GET':
            qs = UnitConversion.objects.filter(
                Q(from_unit=unit) | Q(to_unit=unit)
            )
            serializer = UnitConversionSerializer(qs, many=True, context={'request': request})
            return Response(serializer.data)

        elif request.method == 'POST':
            data = request.data.copy()
            data['from_unit'] = str(unit.id)
            serializer = UnitConversionSerializer(
                data=data,
                context={'request': request},
            )
            if serializer.is_valid():
                serializer.save(tenant_id=_get_tenant_id(request))
                return Response(serializer.data, status=status.HTTP_201_CREATED)
            return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

        elif request.method == 'DELETE':
            conv_id = request.query_params.get('conversion_id')
            if not conv_id:
                return Response(
                    {'error': 'conversion_id query parameter is required.'},
                    status=status.HTTP_400_BAD_REQUEST,
                )
            try:
                conv = UnitConversion.objects.get(id=conv_id, from_unit=unit)
                conv.delete()
                return Response({'success': True, 'message': 'Conversion deleted.'})
            except UnitConversion.DoesNotExist:
                return Response(
                    {'error': 'Conversion not found.'},
                    status=status.HTTP_404_NOT_FOUND,
                )

    @action(detail=False, methods=['get'], url_path='export')
    def export(self, request):
        """Export units to CSV."""
        tenant_id = _get_tenant_id(request)
        qs = Unit.objects.filter(tenant_id=tenant_id)

        status_filter = request.query_params.get('status')
        if status_filter:
            qs = qs.filter(status=status_filter)

        format = request.query_params.get('format', 'csv').lower()
        if format == 'csv':
            return self._export_csv(qs)
        else:
            return self._export_excel(qs)

    def _export_csv(self, qs):
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(['Unit Code', 'Unit Name', 'Symbol', 'Description', 'Status'])
        for unit in qs:
            writer.writerow([unit.unit_code, unit.unit_name, unit.symbol, unit.description, unit.status])
        output.seek(0)
        response = HttpResponse(output.getvalue(), content_type='text/csv')
        response['Content-Disposition'] = 'attachment; filename="units.csv"'
        return response

    def _export_excel(self, qs):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Units'
        ws.append(['Unit Code', 'Unit Name', 'Symbol', 'Description', 'Status'])
        for unit in qs:
            ws.append([unit.unit_code, unit.unit_name, unit.symbol, unit.description, unit.status])
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="units.xlsx"'
        return response

    @action(detail=False, methods=['post'], url_path='import')
    def import_bulk(self, request):
        """Bulk import units."""
        file = request.FILES.get('file')
        tenant_id = _get_tenant_id(request)
        created = []
        errors = []

        try:
            if file:
                if file.name.endswith('.csv'):
                    reader = csv.DictReader(io.StringIO(file.read().decode('utf-8-sig')))
                    rows = list(reader)
                elif file.name.endswith(('.xlsx', '.xls')):
                    wb = openpyxl.load_workbook(file, read_only=True)
                    ws = wb.active
                    headers = [cell.value for cell in next(ws.iter_rows())]
                    rows = []
                    for row in ws.iter_rows(min_row=2, values_only=True):
                        rows.append(dict(zip(headers, row)))
                else:
                    return Response({'error': 'Unsupported file format.'}, status=status.HTTP_400_BAD_REQUEST)
            else:
                serializer = BulkImportSerializer(data=request.data)
                serializer.is_valid(raise_exception=True)
                rows = serializer.validated_data['items']

            for idx, row in enumerate(rows):
                code = str(row.get('Unit Code', row.get('unit_code', ''))).strip()
                name = str(row.get('Unit Name', row.get('unit_name', ''))).strip()
                if not code or not name:
                    errors.append({'index': idx, 'error': 'Unit Code and Name are required.'})
                    continue
                if Unit.objects.filter(tenant_id=tenant_id, unit_code=code).exists():
                    errors.append({'index': idx, 'error': f"Code '{code}' already exists."})
                    continue
                unit = Unit(
                    tenant_id=tenant_id,
                    unit_code=code,
                    unit_name=name,
                    symbol=str(row.get('Symbol', row.get('symbol', ''))).strip(),
                    description=str(row.get('Description', row.get('description', ''))).strip(),
                    status='ACTIVE',
                )
                unit.save()
                created.append(UnitListSerializer(unit).data)

        except Exception as e:
            return Response({'error': f'Failed to parse file: {str(e)}'}, status=status.HTTP_400_BAD_REQUEST)

        return Response({
            'created_count': len(created),
            'error_count': len(errors),
            'created': created,
            'errors': errors,
        }, status=status.HTTP_201_CREATED)


# ============================================================================
# BRAND VIEWSET
# ============================================================================

class BrandViewSet(viewsets.ModelViewSet):
    filter_backends = [DjangoFilterBackend, filters.SearchFilter, filters.OrderingFilter]
    filterset_fields = {
        'status': ['exact'],
    }
    search_fields = ['brand_code', 'brand_name']
    ordering_fields = ['brand_code', 'brand_name', 'created_at', 'status']
    ordering = ['brand_name']

    def get_serializer_class(self):
        if self.action == 'list':
            return BrandSerializer
        return BrandSerializer

    def get_queryset(self):
        tenant_id = _get_tenant_id(self.request)
        qs = Brand.objects.filter(tenant_id=tenant_id)

        show_archived = self.request.query_params.get('show_archived', 'false') == 'true'
        if not show_archived:
            qs = qs.exclude(status='ARCHIVED')

        return qs

    def get_permissions(self):
        perms_map = {
            'create': [CanCreateBrands],
            'update': [CanEditBrands],
            'partial_update': [CanEditBrands],
            'destroy': [CanDeleteBrands],
        }
        for action, perms in perms_map.items():
            if self.action == action:
                return [p() for p in perms]
        return [CanViewBrands()]

    def perform_create(self, serializer):
        user = self.request.user
        serializer.save(
            tenant_id=_get_tenant_id(self.request),
            created_by=user,
            updated_by=user,
        )

    def perform_update(self, serializer):
        serializer.save(updated_by=self.request.user)

    def perform_destroy(self, instance):
        instance.status = 'ARCHIVED'
        instance.save()

    @action(detail=True, methods=['post'])
    def archive(self, request, pk=None):
        """Archive a brand."""
        brand = self.get_object()
        brand.status = 'ARCHIVED'
        brand.save()
        return Response({
            'success': True,
            'message': f"Brand '{brand.brand_name}' archived."
        })

    @action(detail=True, methods=['post'])
    def restore(self, request, pk=None):
        """Restore an archived brand."""
        brand = self.get_object()
        brand.status = 'ACTIVE'
        brand.save()
        return Response({
            'success': True,
            'message': f"Brand '{brand.brand_name}' restored."
        })

    @action(detail=False, methods=['get'], url_path='export')
    def export(self, request):
        """Export brands to CSV."""
        tenant_id = _get_tenant_id(request)
        qs = Brand.objects.filter(tenant_id=tenant_id)

        status_filter = request.query_params.get('status')
        if status_filter:
            qs = qs.filter(status=status_filter)

        format = request.query_params.get('format', 'csv').lower()
        if format == 'csv':
            return self._export_csv(qs)
        else:
            return self._export_excel(qs)

    def _export_csv(self, qs):
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(['Brand Code', 'Brand Name', 'Description', 'Website', 'Status'])
        for brand in qs:
            writer.writerow([
                brand.brand_code, brand.brand_name, brand.description,
                brand.website, brand.status,
            ])
        output.seek(0)
        response = HttpResponse(output.getvalue(), content_type='text/csv')
        response['Content-Disposition'] = 'attachment; filename="brands.csv"'
        return response

    def _export_excel(self, qs):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Brands'
        ws.append(['Brand Code', 'Brand Name', 'Description', 'Website', 'Status'])
        for brand in qs:
            ws.append([
                brand.brand_code, brand.brand_name, brand.description,
                brand.website, brand.status,
            ])
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="brands.xlsx"'
        return response

    @action(detail=False, methods=['post'], url_path='import')
    def import_bulk(self, request):
        """Bulk import brands."""
        file = request.FILES.get('file')
        tenant_id = _get_tenant_id(request)
        created = []
        errors = []

        try:
            if file:
                if file.name.endswith('.csv'):
                    reader = csv.DictReader(io.StringIO(file.read().decode('utf-8-sig')))
                    rows = list(reader)
                elif file.name.endswith(('.xlsx', '.xls')):
                    wb = openpyxl.load_workbook(file, read_only=True)
                    ws = wb.active
                    headers = [cell.value for cell in next(ws.iter_rows())]
                    rows = []
                    for row in ws.iter_rows(min_row=2, values_only=True):
                        rows.append(dict(zip(headers, row)))
                else:
                    return Response({'error': 'Unsupported file format.'}, status=status.HTTP_400_BAD_REQUEST)
            else:
                serializer = BulkImportSerializer(data=request.data)
                serializer.is_valid(raise_exception=True)
                rows = serializer.validated_data['items']

            for idx, row in enumerate(rows):
                code = str(row.get('Brand Code', row.get('brand_code', ''))).strip()
                name = str(row.get('Brand Name', row.get('brand_name', ''))).strip()
                if not code or not name:
                    errors.append({'index': idx, 'error': 'Brand Code and Name are required.'})
                    continue
                if Brand.objects.filter(tenant_id=tenant_id, brand_code=code).exists():
                    errors.append({'index': idx, 'error': f"Code '{code}' already exists."})
                    continue
                brand = Brand(
                    tenant_id=tenant_id,
                    brand_code=code,
                    brand_name=name,
                    description=str(row.get('Description', row.get('description', ''))).strip(),
                    website=str(row.get('Website', row.get('website', ''))).strip(),
                    status='ACTIVE',
                )
                brand.save()
                created.append(BrandListSerializer(brand).data)

        except Exception as e:
            return Response({'error': f'Failed to parse file: {str(e)}'}, status=status.HTTP_400_BAD_REQUEST)

        return Response({
            'created_count': len(created),
            'error_count': len(errors),
            'created': created,
            'errors': errors,
        }, status=status.HTTP_201_CREATED)


# ============================================================================
# INVENTORY ITEM VIEWSET
# ============================================================================

class InventoryItemViewSet(viewsets.ModelViewSet):
    filter_backends = [DjangoFilterBackend, filters.SearchFilter, filters.OrderingFilter]
    filterset_fields = {
        'status': ['exact'],
        'category': ['exact'],
        'brand': ['exact'],
    }
    search_fields = ['item_code', 'item_name', 'description', 'tags']
    ordering_fields = ['item_code', 'item_name', 'created_at', 'updated_at', 'status']
    ordering = ['-created_at']

    def get_serializer_class(self):
        if self.action == 'list':
            return InventoryItemListSerializer
        return InventoryItemSerializer

    def get_queryset(self):
        user = self.request.user
        tenant_id = user.organization_id
        qs = InventoryItem.objects.filter(tenant_id=tenant_id)

        qs = qs.select_related('category', 'unit', 'brand')

        # Filter out archived items from dropdowns (list view based on query param)
        show_archived = self.request.query_params.get('show_archived', 'false') == 'true'
        if not show_archived:
            qs = qs.exclude(status='ARCHIVED')

        return qs

    def get_permissions(self):
        if self.action == 'create':
            permission_classes = [CanCreateItems]
        elif self.action in ['update', 'partial_update']:
            permission_classes = [CanEditItems]
        elif self.action == 'destroy':
            permission_classes = [CanDeleteItems]
        elif self.action == 'import_bulk':
            permission_classes = [CanImportItems]
        elif self.action == 'export':
            permission_classes = [CanExportItems]
        else:
            permission_classes = [CanViewItems]
        return [p() for p in permission_classes]

    def perform_create(self, serializer):
        user = self.request.user
        tenant_id = user.organization_id
        serializer.save(tenant_id=tenant_id, created_by=user, updated_by=user)

    def perform_update(self, serializer):
        serializer.save(updated_by=self.request.user)

    def perform_destroy(self, instance):
        instance.status = 'ARCHIVED'
        instance.save()

    @action(detail=True, methods=['post'])
    def clone(self, request, pk=None):
        """Duplicate an item with a new code/name."""
        original = self.get_object()
        new_code = request.data.get(
            'item_code',
            f"{original.item_code}-CLONE"
        )
        new_name = request.data.get(
            'item_name',
            f"{original.item_name} - New Variant"
        )

        # Check code uniqueness
        if InventoryItem.objects.filter(
            tenant_id=request.user.organization_id,
            item_code=new_code
        ).exists():
            return Response(
                {'error': f"Item code '{new_code}' already exists."},
                status=status.HTTP_400_BAD_REQUEST
            )

        cloned = InventoryItem.objects.create(
            tenant_id=request.user.organization_id,
            item_code=new_code,
            item_name=new_name,
            category=original.category,
            sub_category=original.sub_category,
            unit=original.unit,
            brand=original.brand,
            description=original.description,
            status='ACTIVE',
            custom_fields=original.custom_fields,
            tags=original.tags,
            cloned_from=original,
        )

        serializer = self.get_serializer(cloned)
        return Response(serializer.data, status=status.HTTP_201_CREATED)

    @action(detail=True, methods=['post'])
    def archive(self, request, pk=None):
        """Archive an item."""
        item = self.get_object()
        item.status = 'ARCHIVED'
        item.save()
        return Response({'success': True, 'message': f"Item '{item.item_name}' archived."})

    @action(detail=True, methods=['post'])
    def restore(self, request, pk=None):
        """Restore an archived item."""
        item = self.get_object()
        item.status = 'ACTIVE'
        item.save()
        return Response({'success': True, 'message': f"Item '{item.item_name}' restored."})

    @action(detail=False, methods=['post'], url_path='import')
    def import_bulk(self, request):
        """Bulk import items from JSON payload or Excel file."""
        # Handle file upload
        file = request.FILES.get('file')
        if file:
            return self._import_from_file(file, request)

        # Handle JSON payload
        serializer = BulkImportSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        items_data = serializer.validated_data['items']

        return self._bulk_create_items(items_data, request)

    def _import_from_file(self, file, request):
        """Parse Excel/CSV file and create items."""
        tenant_id = request.user.organization_id
        created = []
        errors = []

        try:
            if file.name.endswith('.csv'):
                reader = csv.DictReader(io.StringIO(file.read().decode('utf-8-sig')))
                rows = list(reader)
            elif file.name.endswith(('.xlsx', '.xls')):
                wb = openpyxl.load_workbook(file, read_only=True)
                ws = wb.active
                headers = [cell.value for cell in next(ws.iter_rows())]
                rows = []
                for row in ws.iter_rows(min_row=2, values_only=True):
                    rows.append(dict(zip(headers, row)))
            else:
                return Response(
                    {'error': 'Unsupported file format. Use CSV or Excel.'},
                    status=status.HTTP_400_BAD_REQUEST
                )

            for idx, row in enumerate(rows):
                item_code = str(row.get('Item Code', row.get('item_code', ''))).strip()
                item_name = str(row.get('Item Name', row.get('item_name', ''))).strip()
                category_name = str(row.get('Category', row.get('category', ''))).strip()
                unit_name = str(row.get('Unit', row.get('unit', ''))).strip()
                brand_name = str(row.get('Brand', row.get('brand', ''))).strip()
                description = str(row.get('Description', row.get('description', ''))).strip()

                if not item_code or not item_name:
                    errors.append({'index': idx, 'error': 'Item Code and Item Name are required.'})
                    continue

                if InventoryItem.objects.filter(
                    tenant_id=tenant_id,
                    item_code=item_code
                ).exists():
                    errors.append({'index': idx, 'error': f"Item code '{item_code}' already exists."})
                    continue

                # Resolve references
                category = None
                if category_name:
                    category = ItemCategory.objects.filter(
                        category_name__iexact=category_name
                    ).first()

                unit = None
                if unit_name:
                    unit = Unit.objects.filter(
                        Q(unit_name__iexact=unit_name) |
                        Q(symbol__iexact=unit_name)
                    ).first()

                brand = None
                if brand_name:
                    brand = Brand.objects.filter(brand_name__iexact=brand_name).first()

                item = InventoryItem(
                    tenant_id=tenant_id,
                    item_code=item_code,
                    item_name=item_name,
                    category=category,
                    unit=unit,
                    brand=brand,
                    description=description,
                    imported_from=file.name,
                    status='ACTIVE',
                )
                item.save()
                created.append(InventoryItemListSerializer(item).data)

        except Exception as e:
            return Response(
                {'error': f'Failed to parse file: {str(e)}'},
                status=status.HTTP_400_BAD_REQUEST
            )

        return Response({
            'created_count': len(created),
            'error_count': len(errors),
            'created': created,
            'errors': errors,
        }, status=status.HTTP_201_CREATED)

    def _bulk_create_items(self, items_data, request):
        """Create multiple items from validated data."""
        tenant_id = request.user.organization_id
        created = []
        errors = []

        for idx, item_data in enumerate(items_data):
            item_code = item_data.get('item_code', '').strip()
            item_name = item_data.get('item_name', '').strip()

            if InventoryItem.objects.filter(
                tenant_id=tenant_id,
                item_code=item_code
            ).exists():
                errors.append({
                    'index': idx,
                    'item_code': item_code,
                    'error': f"Item code '{item_code}' already exists."
                })
                continue

            category = None
            cat_name = item_data.get('category', '')
            if cat_name:
                category = ItemCategory.objects.filter(category_name__iexact=cat_name).first()

            unit = None
            unit_name = item_data.get('unit', '')
            if unit_name:
                unit = Unit.objects.filter(
                    Q(unit_name__iexact=unit_name) |
                    Q(symbol__iexact=unit_name)
                ).first()

            brand = None
            brand_name = item_data.get('brand', '')
            if brand_name:
                brand = Brand.objects.filter(brand_name__iexact=brand_name).first()

            item = InventoryItem(
                tenant_id=tenant_id,
                item_code=item_code,
                item_name=item_name,
                category=category,
                unit=unit,
                brand=brand,
                description=item_data.get('description', ''),
                status='ACTIVE',
            )
            item.save()
            created.append(InventoryItemListSerializer(item).data)

        return Response({
            'created_count': len(created),
            'error_count': len(errors),
            'created': created,
            'errors': errors,
        }, status=status.HTTP_201_CREATED)

    @action(detail=False, methods=['get'], url_path='export')
    def export(self, request):
        """Export items to CSV."""
        tenant_id = request.user.organization_id
        items = InventoryItem.objects.filter(tenant_id=tenant_id).select_related(
            'category', 'unit', 'brand'
        )

        # Apply filters
        status_filter = request.query_params.get('status')
        if status_filter:
            items = items.filter(status=status_filter)

        category_filter = request.query_params.get('category')
        if category_filter:
            items = items.filter(category_id=category_filter)

        format = request.query_params.get('format', 'csv').lower()

        if format == 'csv':
            return self._export_csv(items)
        else:
            return self._export_excel(items)

    def _export_csv(self, items):
        """Export items as CSV."""
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow([
            'Item Code', 'Item Name', 'Category', 'Sub Category',
            'Unit', 'Brand', 'Description', 'Status', 'Tags'
        ])

        for item in items:
            writer.writerow([
                item.item_code,
                item.item_name,
                item.category.category_name if item.category else '',
                item.sub_category,
                item.unit.unit_name if item.unit else '',
                item.brand.brand_name if item.brand else '',
                item.description,
                item.status,
                ', '.join(item.tags) if item.tags else '',
            ])

        output.seek(0)
        response = HttpResponse(
            output.getvalue(),
            content_type='text/csv'
        )
        response['Content-Disposition'] = 'attachment; filename="inventory_items.csv"'
        return response

    def _export_excel(self, items):
        """Export items as Excel (.xlsx)."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Inventory Items'
        ws.append([
            'Item Code', 'Item Name', 'Category', 'Sub Category',
            'Unit', 'Brand', 'Description', 'Status', 'Tags'
        ])

        for item in items:
            ws.append([
                item.item_code,
                item.item_name,
                item.category.category_name if item.category else '',
                item.sub_category,
                item.unit.unit_name if item.unit else '',
                item.brand.brand_name if item.brand else '',
                item.description,
                item.status,
                ', '.join(item.tags) if item.tags else '',
            ])

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="inventory_items.xlsx"'
        return response


# ============================================================================
# CUSTOM FIELD DEFINITION VIEWSET
# ============================================================================

class CustomFieldDefinitionViewSet(viewsets.ModelViewSet):
    queryset = CustomFieldDefinition.objects.filter(is_active=True)
    serializer_class = CustomFieldDefinitionSerializer
    permission_classes = [CanViewItems]
    filter_backends = [filters.SearchFilter]
    search_fields = ['field_name', 'field_label']

    def get_permissions(self):
        if self.action in ['create', 'update', 'partial_update', 'destroy']:
            permission_classes = [CanCreateItems]
        else:
            permission_classes = [CanViewItems]
        return [p() for p in permission_classes]


# ============================================================================
# STOCK LEDGER VIEWSET
# ============================================================================

class StockLedgerViewSet(viewsets.ReadOnlyModelViewSet):
    """View stock ledger entries."""
    serializer_class = StockLedgerSerializer
    filter_backends = [DjangoFilterBackend, filters.SearchFilter, filters.OrderingFilter]
    filterset_fields = {
        'item': ['exact'],
        'location': ['exact'],
        'transaction_type': ['exact'],
        'created_at': ['gte', 'lte', 'date'],
    }
    search_fields = ['item__item_code', 'item__item_name', 'description', 'reference_id']
    ordering_fields = ['created_at', 'transaction_type', 'quantity']
    ordering = ['-created_at']

    def get_queryset(self):
        return StockLedger.objects.filter(
            tenant_id=_get_tenant_id(self.request)
        ).select_related('item', 'location')

    def get_permissions(self):
        return [CanViewStock()]


# ============================================================================
# STOCK LEDGER ENTRY CREATE (POST only)
# ============================================================================

class StockLedgerEntryView(viewsets.ViewSet):
    """Create stock ledger entries."""
    permission_classes = [CanCreateItems]

    def create(self, request):
        serializer = CreateLedgerEntrySerializer(data=request.data)
        serializer.is_valid(raise_exception=True)

        data = serializer.validated_data
        entry = create_ledger_entry(
            tenant_id=_get_tenant_id(request),
            item_id=data['item_id'],
            transaction_type=data['transaction_type'],
            quantity=data['quantity'],
            location_id=data.get('location_id'),
            unit_cost=data.get('unit_cost'),
            reference_type=data.get('reference_type', ''),
            reference_id=data.get('reference_id', ''),
            description=data.get('description', ''),
            created_by=request.user,
        )

        return Response(
            StockLedgerSerializer(entry, context={'request': request}).data,
            status=status.HTTP_201_CREATED,
        )


# ============================================================================
# STOCK AVAILABILITY VIEW
# ============================================================================

class StockAvailabilityViewSet(viewsets.ViewSet):
    """
    View stock availability — calculated from the ledger in real-time.
    Never reads stored quantities.
    """
    permission_classes = [CanViewStock]

    def get_permissions(self):
        if self.action in ['export']:
            return [CanExportStock()]
        elif self.action in ['snapshot']:
            return [CanViewStockSnapshot()]
        elif self.action in ['valuation']:
            return [CanViewStockValuation()]
        return [CanViewStock()]

    def list(self, request):
        """GET /api/inventory/stock/availability/ — all items with stock"""
        tenant_id = _get_tenant_id(request)
        filters = {
            'search': request.query_params.get('search', ''),
            'category': request.query_params.get('category'),
            'brand': request.query_params.get('brand'),
            'status': request.query_params.get('status'),
        }
        # Remove empty filters
        filters = {k: v for k, v in filters.items() if v}

        results = get_all_availability(tenant_id, filters)

        # In-memory pagination
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 50))
        start = (page - 1) * page_size
        end = start + page_size
        page_data = results[start:end]

        return Response({
            'count': len(results),
            'results': page_data,
        })

    def retrieve(self, request, pk=None):
        """GET /api/inventory/stock/availability/{item_id}/ — item detail"""
        tenant_id = _get_tenant_id(request)
        data = get_item_availability(pk, tenant_id)
        return Response(data)

    @action(detail=False, methods=['get'], url_path='location/(?P<location_id>[^/.]+)')
    def by_location(self, request, location_id=None):
        """GET /api/inventory/stock/availability/location/{location_id}/"""
        tenant_id = _get_tenant_id(request)
        data = get_location_availability(location_id, tenant_id)
        return Response(data)

    @action(detail=False, methods=['get'], url_path='low-stock')
    def low_stock(self, request):
        """GET /api/inventory/stock/availability/low-stock/"""
        tenant_id = _get_tenant_id(request)
        data = get_low_stock_items(tenant_id)
        return Response({'count': len(data), 'results': data})

    @action(detail=False, methods=['get'], url_path='out-of-stock')
    def out_of_stock(self, request):
        """GET /api/inventory/stock/availability/out-of-stock/"""
        tenant_id = _get_tenant_id(request)
        data = get_out_of_stock_items(tenant_id)
        return Response({'count': len(data), 'results': data})

    @action(detail=False, methods=['get'], url_path='valuation')
    def valuation(self, request):
        """GET /api/inventory/stock/availability/valuation/"""
        tenant_id = _get_tenant_id(request)
        filters = {
            'category': request.query_params.get('category'),
            'brand': request.query_params.get('brand'),
        }
        filters = {k: v for k, v in filters.items() if v}
        data = get_valuation(tenant_id, filters)
        return Response(data)

    @action(detail=False, methods=['get'], url_path='snapshot')
    def snapshot(self, request):
        """GET /api/inventory/stock/availability/snapshot/?as_of=2025-12-31"""
        tenant_id = _get_tenant_id(request)
        as_of = request.query_params.get('as_of')
        if as_of:
            try:
                as_of_date = datetime.strptime(as_of, '%Y-%m-%d').date()
            except ValueError:
                return Response(
                    {'error': 'Invalid date format. Use YYYY-MM-DD.'},
                    status=status.HTTP_400_BAD_REQUEST,
                )
        else:
            as_of_date = timezone.now().date()

        data = get_snapshot(tenant_id, as_of_date)
        return Response(data)

    @action(detail=False, methods=['get'], url_path='export')
    def export(self, request):
        """GET /api/inventory/stock/availability/export/?format=csv"""
        tenant_id = _get_tenant_id(request)
        results = get_all_availability(tenant_id)
        fmt = request.query_params.get('format', 'csv').lower()

        if fmt == 'csv':
            output = io.StringIO()
            writer = csv.writer(output)
            writer.writerow([
                'Item Code', 'Item Name', 'Category', 'Brand',
                'Physical', 'Reserved', 'Available',
                'In Transit', 'Damaged',
                'Cost Price', 'Selling Price',
                'Cost Value', 'Selling Value',
            ])
            for r in results:
                writer.writerow([
                    r['item_code'], r['item_name'],
                    r['category_name'], r['brand_name'],
                    r['physical'], r['reserved'], r['available'],
                    r['in_transit'], r['damaged'],
                    r['cost_price'], r['selling_price'],
                    r['cost_value'], r['selling_value'],
                ])
            output.seek(0)
            response = HttpResponse(output.getvalue(), content_type='text/csv')
            response['Content-Disposition'] = 'attachment; filename="stock_availability.csv"'
            return response
        else:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = 'Stock Availability'
            ws.append([
                'Item Code', 'Item Name', 'Category', 'Brand',
                'Physical', 'Reserved', 'Available',
                'In Transit', 'Damaged',
                'Cost Price', 'Selling Price',
                'Cost Value', 'Selling Value',
            ])
            for r in results:
                ws.append([
                    r['item_code'], r['item_name'],
                    r['category_name'], r['brand_name'],
                    float(r['physical']), float(r['reserved']),
                    float(r['available']), float(r['in_transit']),
                    float(r['damaged']),
                    float(r['cost_price']) if r['cost_price'] else '',
                    float(r['selling_price']) if r['selling_price'] else '',
                    float(r['cost_value']), float(r['selling_value']),
                ])
            output = io.BytesIO()
            wb.save(output)
            output.seek(0)
            response = HttpResponse(
                output.getvalue(),
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
            response['Content-Disposition'] = 'attachment; filename="stock_availability.xlsx"'
            return response


# ============================================================================
# LOCATION TYPE VIEWSET
# ============================================================================

class LocationTypeViewSet(viewsets.ModelViewSet):
    filter_backends = [DjangoFilterBackend, filters.SearchFilter, filters.OrderingFilter]
    filterset_fields = {
        'status': ['exact'],
    }
    search_fields = ['type_code', 'type_name']
    ordering_fields = ['type_code', 'type_name', 'created_at', 'status']
    ordering = ['type_name']

    def get_serializer_class(self):
        if self.action == 'list':
            return LocationTypeSerializer
        return LocationTypeSerializer

    def get_queryset(self):
        tenant_id = _get_tenant_id(self.request)
        qs = InventoryLocationType.objects.filter(tenant_id=tenant_id)
        return qs

    def get_permissions(self):
        perms_map = {
            'create': [CanCreateLocationTypes],
            'update': [CanEditLocationTypes],
            'partial_update': [CanEditLocationTypes],
            'destroy': [CanDeleteLocationTypes],
        }
        for action, perms in perms_map.items():
            if self.action == action:
                return [p() for p in perms]
        return [CanViewLocationTypes()]

    def perform_create(self, serializer):
        serializer.save(tenant_id=_get_tenant_id(self.request))

    def perform_destroy(self, instance):
        if instance.locations.exists():
            raise serializers.ValidationError(
                "Cannot delete location type in use by existing locations."
            )
        instance.delete()


# ============================================================================
# TRANSFER VIEWSET
# ============================================================================

class TransferViewSet(viewsets.ModelViewSet):
    """Stock transfers with full workflow: submit → approve → dispatch → receive"""
    filter_backends = [DjangoFilterBackend, filters.SearchFilter, filters.OrderingFilter]
    filterset_fields = {
        'status': ['exact'],
        'source_location': ['exact'],
        'destination_location': ['exact'],
        'transfer_type': ['exact'],
        'transfer_date': ['gte', 'lte'],
    }
    search_fields = ['transfer_number', 'remarks']
    ordering_fields = ['transfer_number', 'transfer_date', 'created_at', 'status']
    ordering = ['-created_at']

    def get_serializer_class(self):
        if self.action == 'list':
            return TransferListSerializer
        elif self.action in ['create']:
            return CreateTransferSerializer
        elif self.action in ['receive']:
            return ReceiveTransferSerializer
        return TransferDetailSerializer

    def get_queryset(self):
        tenant_id = _get_tenant_id(self.request)
        return InventoryTransfer.objects.filter(tenant_id=tenant_id).select_related(
            'source_location', 'destination_location', 'created_by'
        ).prefetch_related('items__item')

    def get_permissions(self):
        perms_map = {
            'create': [CanCreateTransfers],
            'update': [CanEditTransfers],
            'partial_update': [CanEditTransfers],
            'destroy': [CanDeleteTransfers],
            'submit': [CanSubmitTransfer],
            'approve': [CanApproveTransfer],
            'reject': [CanApproveTransfer],
            'dispatch': [CanApproveTransfer],
            'receive': [CanReceiveTransfer],
            'export': [CanExportTransfers],
        }
        for action, perms in perms_map.items():
            if self.action == action:
                return [p() for p in perms]
        return [CanViewTransfers()]

    def create(self, request, *args, **kwargs):
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        tenant_id = _get_tenant_id(request)
        try:
            transfer = create_transfer(tenant_id, serializer.validated_data, request.user)
        except ValueError as e:
            return Response({'error': str(e)}, status=status.HTTP_400_BAD_REQUEST)
        output_serializer = TransferDetailSerializer(
            transfer, context={'request': request}
        )
        return Response(output_serializer.data, status=status.HTTP_201_CREATED)

    @action(detail=True, methods=['post'])
    def submit(self, request, pk=None):
        """Submit transfer for approval."""
        transfer = self.get_object()
        try:
            transfer = submit_transfer(transfer, request.user)
        except ValueError as e:
            return Response({'error': str(e)}, status=status.HTTP_400_BAD_REQUEST)
        serializer = TransferDetailSerializer(transfer, context={'request': request})
        return Response(serializer.data)

    @action(detail=True, methods=['post'])
    def approve(self, request, pk=None):
        """Approve a transfer."""
        transfer = self.get_object()
        notes = request.data.get('notes', '')
        try:
            transfer = approve_transfer(transfer, request.user, notes)
        except ValueError as e:
            return Response({'error': str(e)}, status=status.HTTP_400_BAD_REQUEST)
        serializer = TransferDetailSerializer(transfer, context={'request': request})
        return Response(serializer.data)

    @action(detail=True, methods=['post'])
    def reject(self, request, pk=None):
        """Reject a transfer."""
        transfer = self.get_object()
        notes = request.data.get('notes', '')
        try:
            transfer = reject_transfer(transfer, request.user, notes)
        except ValueError as e:
            return Response({'error': str(e)}, status=status.HTTP_400_BAD_REQUEST)
        serializer = TransferDetailSerializer(transfer, context={'request': request})
        return Response(serializer.data)

    @action(detail=True, methods=['post'])
    def dispatch(self, request, pk=None):
        """Dispatch a transfer — creates TRANSFER_OUT ledger entries."""
        transfer = self.get_object()
        try:
            transfer = dispatch_transfer(transfer, request.user)
        except ValueError as e:
            return Response({'error': str(e)}, status=status.HTTP_400_BAD_REQUEST)
        serializer = TransferDetailSerializer(transfer, context={'request': request})
        return Response(serializer.data)

    @action(detail=True, methods=['post'])
    def receive(self, request, pk=None):
        """Receive a transfer — creates TRANSFER_IN ledger entries."""
        transfer = self.get_object()
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        try:
            transfer = receive_transfer(
                transfer, request.user,
                serializer.validated_data.get('items')
            )
        except ValueError as e:
            return Response({'error': str(e)}, status=status.HTTP_400_BAD_REQUEST)
        output = TransferDetailSerializer(transfer, context={'request': request})
        return Response(output.data)

    @action(detail=True, methods=['post'])
    def cancel(self, request, pk=None):
        """Cancel a transfer (DRAFT or PENDING_APPROVAL only)."""
        transfer = self.get_object()
        try:
            transfer = cancel_transfer(transfer, request.user)
        except ValueError as e:
            return Response({'error': str(e)}, status=status.HTTP_400_BAD_REQUEST)
        serializer = TransferDetailSerializer(transfer, context={'request': request})
        return Response(serializer.data)

    @action(detail=False, methods=['get'], url_path='export')
    def export(self, request):
        """Export transfers to CSV/Excel."""
        tenant_id = _get_tenant_id(request)
        qs = InventoryTransfer.objects.filter(tenant_id=tenant_id).select_related(
            'source_location', 'destination_location', 'created_by'
        )

        fmt = request.query_params.get('format', 'csv').lower()
        if fmt == 'csv':
            output = io.StringIO()
            writer = csv.writer(output)
            writer.writerow([
                'Transfer Number', 'Date', 'Source', 'Destination',
                'Type', 'Status', 'Items', 'Created By', 'Created At'
            ])
            for t in qs:
                writer.writerow([
                    t.transfer_number, t.transfer_date,
                    t.source_location.location_name,
                    t.destination_location.location_name,
                    t.transfer_type, t.status,
                    t.items.count(),
                    t.created_by.email if t.created_by else '',
                    t.created_at,
                ])
            output.seek(0)
            response = HttpResponse(output.getvalue(), content_type='text/csv')
            response['Content-Disposition'] = 'attachment; filename="transfers.csv"'
            return response
        else:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = 'Transfers'
            ws.append([
                'Transfer Number', 'Date', 'Source', 'Destination',
                'Type', 'Status', 'Items', 'Created By', 'Created At'
            ])
            for t in qs:
                ws.append([
                    t.transfer_number, str(t.transfer_date),
                    t.source_location.location_name,
                    t.destination_location.location_name,
                    t.transfer_type, t.status,
                    t.items.count(),
                    t.created_by.email if t.created_by else '',
                    str(t.created_at),
                ])
            output = io.BytesIO()
            wb.save(output)
            output.seek(0)
            response = HttpResponse(
                output.getvalue(),
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
            response['Content-Disposition'] = 'attachment; filename="transfers.xlsx"'
            return response


# ============================================================================
# LOCATION VIEWSET
# ============================================================================

class LocationViewSet(viewsets.ModelViewSet):
    filter_backends = [DjangoFilterBackend, filters.SearchFilter, filters.OrderingFilter]
    filterset_fields = {
        'status': ['exact'],
        'location_type': ['exact'],
        'city': ['exact', 'icontains'],
        'state': ['exact', 'icontains'],
        'country': ['exact', 'icontains'],
    }
    search_fields = ['location_code', 'location_name', 'city', 'state', 'country']
    ordering_fields = ['location_code', 'location_name', 'city', 'status', 'created_at']
    ordering = ['location_name']

    def get_serializer_class(self):
        if self.action == 'list':
            return LocationListSerializer
        elif self.action == 'tree':
            return LocationTreeSerializer
        return LocationSerializer

    def get_queryset(self):
        tenant_id = _get_tenant_id(self.request)
        qs = InventoryLocation.objects.filter(tenant_id=tenant_id)

        show_archived = self.request.query_params.get('show_archived', 'false') == 'true'
        if not show_archived and self.action != 'tree':
            qs = qs.exclude(status='ARCHIVED')

        return qs

    def get_permissions(self):
        perms_map = {
            'create': [CanCreateLocations],
            'update': [CanEditLocations],
            'partial_update': [CanEditLocations],
            'destroy': [CanDeleteLocations],
        }
        for action, perms in perms_map.items():
            if self.action == action:
                return [p() for p in perms]
        return [CanViewLocations()]

    def perform_create(self, serializer):
        user = self.request.user
        serializer.save(
            tenant_id=_get_tenant_id(self.request),
            created_by=user,
            updated_by=user,
        )

    def perform_update(self, serializer):
        serializer.save(updated_by=self.request.user)

    def perform_destroy(self, instance):
        instance.status = 'ARCHIVED'
        instance.save()

    @action(detail=True, methods=['post'])
    def archive(self, request, pk=None):
        """Archive a location."""
        location = self.get_object()
        location.status = 'ARCHIVED'
        location.save()
        return Response({
            'success': True,
            'message': f"Location '{location.location_name}' archived."
        })

    @action(detail=True, methods=['post'])
    def restore(self, request, pk=None):
        """Restore an archived location."""
        location = self.get_object()
        location.status = 'ACTIVE'
        location.save()
        return Response({
            'success': True,
            'message': f"Location '{location.location_name}' restored."
        })

    @action(detail=False, methods=['get'])
    def tree(self, request):
        """Return hierarchical location tree."""
        tenant_id = _get_tenant_id(request)
        locations = InventoryLocation.objects.filter(
            tenant_id=tenant_id,
            parent_location__isnull=True,
        )
        show_archived = request.query_params.get('show_archived', 'false') == 'true'
        if not show_archived:
            locations = locations.exclude(status='ARCHIVED')

        serializer = LocationTreeSerializer(
            locations, many=True,
            context={'request': request, 'filters': {'status': None}}
        )
        return Response(serializer.data)

    @action(detail=False, methods=['get'], url_path='export')
    def export(self, request):
        """Export locations to CSV."""
        tenant_id = _get_tenant_id(request)
        qs = InventoryLocation.objects.filter(tenant_id=tenant_id)

        status_filter = request.query_params.get('status')
        if status_filter:
            qs = qs.filter(status=status_filter)

        format = request.query_params.get('format', 'csv').lower()
        if format == 'csv':
            return self._export_csv(qs)
        else:
            return self._export_excel(qs)

    def _export_csv(self, qs):
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow([
            'Location Code', 'Location Name', 'Type', 'Parent',
            'City', 'State', 'Country', 'Status'
        ])
        for loc in qs:
            writer.writerow([
                loc.location_code,
                loc.location_name,
                loc.location_type.type_name if loc.location_type else '',
                loc.parent_location.location_code if loc.parent_location else '',
                loc.city, loc.state, loc.country, loc.status,
            ])
        output.seek(0)
        response = HttpResponse(output.getvalue(), content_type='text/csv')
        response['Content-Disposition'] = 'attachment; filename="locations.csv"'
        return response

    def _export_excel(self, qs):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Locations'
        ws.append([
            'Location Code', 'Location Name', 'Type', 'Parent',
            'City', 'State', 'Country', 'Status'
        ])
        for loc in qs:
            ws.append([
                loc.location_code,
                loc.location_name,
                loc.location_type.type_name if loc.location_type else '',
                loc.parent_location.location_code if loc.parent_location else '',
                loc.city, loc.state, loc.country, loc.status,
            ])
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="locations.xlsx"'
        return response

    @action(detail=False, methods=['post'], url_path='import')
    def import_bulk(self, request):
        """Bulk import locations from CSV/Excel or JSON."""
        file = request.FILES.get('file')
        tenant_id = _get_tenant_id(request)
        created = []
        errors = []

        try:
            if file:
                if file.name.endswith('.csv'):
                    reader = csv.DictReader(io.StringIO(file.read().decode('utf-8-sig')))
                    rows = list(reader)
                elif file.name.endswith(('.xlsx', '.xls')):
                    wb = openpyxl.load_workbook(file, read_only=True)
                    ws = wb.active
                    headers = [cell.value for cell in next(ws.iter_rows())]
                    rows = []
                    for row in ws.iter_rows(min_row=2, values_only=True):
                        rows.append(dict(zip(headers, row)))
                else:
                    return Response({'error': 'Unsupported file format.'}, status=status.HTTP_400_BAD_REQUEST)
            else:
                serializer = BulkImportSerializer(data=request.data)
                serializer.is_valid(raise_exception=True)
                rows = serializer.validated_data['items']

            for idx, row in enumerate(rows):
                code = str(row.get('Location Code', row.get('location_code', ''))).strip()
                name = str(row.get('Location Name', row.get('location_name', ''))).strip()
                type_name = str(row.get('Type', row.get('location_type', ''))).strip()

                if not code or not name:
                    errors.append({'index': idx, 'error': 'Location Code and Name are required.'})
                    continue

                if InventoryLocation.objects.filter(tenant_id=tenant_id, location_code=code).exists():
                    errors.append({'index': idx, 'error': f"Code '{code}' already exists."})
                    continue

                location_type = None
                if type_name:
                    location_type = InventoryLocationType.objects.filter(
                        tenant_id=tenant_id,
                        type_name__iexact=type_name
                    ).first()

                parent_code = str(row.get('Parent', row.get('parent_location', ''))).strip()
                parent = None
                if parent_code:
                    parent = InventoryLocation.objects.filter(
                        tenant_id=tenant_id, location_code=parent_code
                    ).first()

                loc = InventoryLocation(
                    tenant_id=tenant_id,
                    location_code=code,
                    location_name=name,
                    location_type=location_type,
                    parent_location=parent,
                    city=str(row.get('City', row.get('city', ''))).strip(),
                    state=str(row.get('State', row.get('state', ''))).strip(),
                    country=str(row.get('Country', row.get('country', ''))).strip(),
                    status='ACTIVE',
                )
                loc.save()
                created.append(LocationListSerializer(loc).data)

        except Exception as e:
            return Response({'error': f'Failed to parse file: {str(e)}'}, status=status.HTTP_400_BAD_REQUEST)

        return Response({
            'created_count': len(created),
            'error_count': len(errors),
            'created': created,
            'errors': errors,
        }, status=status.HTTP_201_CREATED)
